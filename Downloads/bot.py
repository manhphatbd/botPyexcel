import os
import openpyxl
from openpyxl.styles import Font, Alignment
from telegram import Update
from openpyxl.drawing.image import Image
from telegram.ext import Application, CommandHandler, MessageHandler, filters, ContextTypes
from pydrive.auth import GoogleAuth
from pydrive.drive import GoogleDrive
from vietnam_number import n2w

# ====== CONFIG ======
TOKEN = "your token"
BASE_FILE = "baogia.xlsx"
MARKER = "###DATA_END###"


# ====== Utils ======
def get_new_filename(base_name="ketqua.xlsx"):
    if not os.path.exists(base_name):
        return base_name
    i = 1
    while True:
        new_name = f"ketqua-{i}.xlsx"
        if not os.path.exists(new_name):
            return new_name
        i += 1


def safe_int(value):
    try:
        return int(str(value).replace(".", "").replace(",", "").strip())
    except:
        return 0


# ====== Parse input ======
def parse_input(text):
    # Tách các sản phẩm từ dấu chấm (.), loại bỏ khoảng trắng thừa
    products = [p.strip() for p in text.split(".") if p.strip()]
    parsed = []

    for prod in products:
        # Tách các trường trong mỗi sản phẩm bằng dấu phẩy (",")
        fields = [f.strip() for f in prod.split(",")]

        # Kiểm tra nếu có ít hơn 8 trường (không hợp lệ, bỏ qua sản phẩm)
        if len(fields) < 8:
            continue

        # Lấy 8 trường đầu tiên
        fields = fields[:8]

        # Thêm sản phẩm vào danh sách parsed
        parsed.append({
            "stt": fields[0],  # Trường đầu tiên là STT
            "ten": fields[1],  # Tên sản phẩm
            "dvt": fields[2],  # Đơn vị tính
            "ma": fields[3],   # Mã sản phẩm
            "xuatxu": fields[4],  # Xuất xứ
            "baohanh": fields[5],  # Bảo hành
            "soluong": safe_int(fields[6]),  # Số lượng (ép kiểu int)
            "dongia": safe_int(fields[7])   # Đơn giá (ép kiểu int)
        })

    return parsed


# ====== Google Drive ======
def upload_to_drive(file_path):
    gauth = GoogleAuth()
    gauth.LoadCredentialsFile("credentials.json")

    if gauth.credentials is None:
        gauth.LocalWebserverAuth()
    elif gauth.access_token_expired:
        gauth.Refresh()
    else:
        gauth.Authorize()

    gauth.SaveCredentialsFile("credentials.json")

    drive = GoogleDrive(gauth)
    file = drive.CreateFile({'title': os.path.basename(file_path)})
    file.SetContentFile(file_path)
    file.Upload()

    return file['id']

def to_vietnamese_currency(number):
    number = safe_int(number)  # chuẩn hóa thành int sạch
    if not number:
        return "Bằng chữ: Không đồng."
    # n2w cần chuỗi, không phải int
    text = n2w(str(number))
    text = text.capitalize()
    return f"Bằng chữ: {text} chẵn."

# ====== Excel xử lý ======
def update_excel(file_path, input_text):
    data = parse_input(input_text)

    if not data:
        raise Exception("Không có dữ liệu hợp lệ!")

    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active

    start_row = 11

    # ✅ Tìm marker
    marker_row = None
    for row in range(start_row, sheet.max_row + 1):
        val = sheet.cell(row=row, column=1).value
        if val and MARKER in str(val):
            marker_row = row
            break

    if not marker_row:
        raise Exception("Không tìm thấy marker ###DATA_END### trong file Excel!")
# ✅ Chỉ unmerge vùng dữ liệu (từ row 10 trở xuống), giữ nguyên header
    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.min_row >= 10 and merged_range.max_row <= marker_row:
            sheet.unmerge_cells(str(merged_range))

    # ✅ Xóa dữ liệu cũ (chỉ xóa phía trên marker)
    #sheet.delete_rows(start_row, marker_row - start_row)
    # Clear dữ liệu cũ (không xóa row)
    #for row in range(start_row, marker_row):
     #   for col in range(1, 10):
      #      sheet.cell(row=row, column=col).value = None
    # ✅ Xóa dữ liệu cũ mà không làm mất định dạng
    for row in range(start_row, marker_row):
        for col in range(1, 10):  # Xóa giá trị trong phạm vi cột 1 → 9
            cell = sheet.cell(row=row, column=col)
            cell.value = None  # Chỉ xóa giá trị, không thay đổi định dạng

    # ✅ Ghi header
    headers = ["STT","Tên hạng mục","Đvt","Mã hiệu/Mô tả","Xuất xứ","Bảo hành","Số lượng","Đơn giá","Thành tiền"]
    for col, h in enumerate(headers, start=1):
        sheet.cell(row=10, column=col, value=h)

    # ✅ Ghi dữ liệu
    for idx, item in enumerate(data, start=1):
        row = start_row + idx - 1

        sheet.cell(row=row, column=1, value=idx)
        sheet.cell(row=row, column=2, value=item["ten"])
        sheet.cell(row=row, column=3, value=item["dvt"])
        sheet.cell(row=row, column=4, value=item["ma"])
        sheet.cell(row=row, column=5, value=item["xuatxu"])
        sheet.cell(row=row, column=6, value=item["baohanh"])
        sheet.cell(row=row, column=7, value=item["soluong"])
        sheet.cell(row=row, column=8, value=item["dongia"])
        sheet.cell(row=row, column=9, value=item["soluong"] * item["dongia"])

    # ✅ Tổng tiền
    total_row = start_row + len(data)
    total_amount = sum(i["soluong"] * i["dongia"] for i in data)

# Gộp từ cột 1 → 8
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=8)

    cell_total = sheet.cell(row=total_row, column=1)
    cell_total.value = "Tổng thanh toán"
    cell_total.font = Font(bold=True)
    cell_total.alignment = Alignment(horizontal="center", vertical="center")

# Số tiền ở cột 9
    cell_money = sheet.cell(row=total_row, column=9)
    cell_money.value = total_amount
    cell_money.font = Font(bold=True)

    row_text = total_row + 1

# Merge 10 cột (1 → 10)
    sheet.merge_cells(start_row=row_text, start_column=1, end_row=row_text, end_column=10)

    cell2 = sheet.cell(row=row_text, column=1)
    cell2.value = to_vietnamese_currency(total_amount)

    cell2.font = Font(bold=True, italic=True)
    cell2.alignment = Alignment(horizontal="center", vertical="center")
    sheet.delete_rows(row_text + 1, marker_row - (row_text + 1) + 1)
    # ✅ Chèn lại logo
    #logo_path = "logo.png"  # đường dẫn file logo của bạn
    #if os.path.exists(logo_path):
     #   img = Image(logo_path)
        # Đặt kích thước theo inch (Excel ~96 dpi)
      #  img.width = int(1.51 * 96)   # ~145 px
       # img.height = int(1.46 * 96)  # ~140 px

        # Vị trí chèn (ví dụ A1)
        #sheet.add_image(img, "B3")

    output_path = get_new_filename()
    wb.save(output_path)

    file_id = upload_to_drive(output_path)

    return output_path, file_id


# ====== Telegram Bot ======
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Gửi dữ liệu dạng:\n"
        "1, Laptop Dell, cái, DELL123, USA, 12 tháng, 2, 15000000 . ..."
    )


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        input_text = update.message.text

        output_file, file_id = update_excel(BASE_FILE, input_text)

        link = f"https://drive.google.com/file/d/{file_id}/view"

        await update.message.reply_document(document=open(output_file, "rb"))
        await update.message.reply_text(f"Link Drive: {link}")

    except Exception as e:
        await update.message.reply_text(f"Lỗi: {str(e)}")


def main():
    if not TOKEN:
        raise Exception("Bạn chưa set TELEGRAM_TOKEN")

    app = Application.builder().token(TOKEN).build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))

    print("Bot đang chạy...")
    app.run_polling()


if __name__ == "__main__":
# Input mẫu: mỗi sản phẩm cách nhau bằng dấu chấm ".", mỗi trường cách nhau bằng dấu phẩy ","
    sample_input = "1, bàn epoxy vuông, cái, BVE-X, Việt Nam, 5 năm, 2, 16000000. 2, Bàn epoxy tròng, cái, BTE-E, China, 3 năm, 5, 600000. 3, Ruột gà, Cuộn, KEY789, Vietnam, 12 tháng, 7, 1200000. 5, HikVision, Con, PY93, Việt Nam, 2 tháng, 8, 860000. 6, Tapo C222, Con, LOGI456, China, 24 tháng, 5, 580000. 2, Imou, Con, BKVision, China, 24 tháng, 5, 385000. 6, Thiết kế webiste, Page, TL43299, Việt Nam, 12 tháng, , 12000000. 2, Software, Slot, TL9989, Việt Nam, Vĩnh viễn, 1, 25000000."
    output_file = update_excel("baogia.xlsx", sample_input)
    print(f"File Excel đã được tạo: {output_file}")
    
    main()
